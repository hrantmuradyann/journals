#!/usr/bin/env python3
"""Read news_articles.xlsx and sort the articles using plain keyword rules.

No API or machine learning: every decision here is a list of words you can edit.

- TOPICS        -> each article gets exactly one topic (highest keyword score wins)
- EDUCATIONAL   -> title-only patterns for "explainer" style articles (a tag, on top of the topic)
- REGIONS       -> region tags such as Asia (a tag, on top of the topic)
- same-story    -> articles from different publications whose titles share rare words
"""

import collections
import os
import re
from datetime import datetime

from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Keyword lists (edit freely). Matching is case-insensitive, ignores hyphens
# and dots inside abbreviations ("A.I." == "ai", "data-centre" == "data centre"),
# and allows common endings: s, es, ed, ing, er, ers, ist, al, ic, ian.
# ---------------------------------------------------------------------------

TOPICS = [
    ("Technology & AI", [
        "ai", "artificial intelligence", "chatgpt", "openai", "anthropic", "claude", "gemini",
        "llm", "large language model", "machine learning", "chip", "semiconductor", "nvidia",
        "software", "app", "startup", "start up", "robot", "quantum", "cyber", "hack", "hacker",
        "data centre", "data center", "cloud", "smartphone", "iphone", "google", "alphabet",
        "apple", "microsoft", "meta", "amazon", "tesla", "spacex", "tech", "technology",
        "algorithm", "crypto", "bitcoin", "stablecoin", "silicon valley", "tiktok", "social media",
        "automation", "self driving", "autonomous", "drone", "satellite", "internet", "online",
        "digital", "computer", "computing", "video game", "streaming", "platform", "big tech",
        "coding", "battery", "electric vehicle", "rocket", "space", "huawei", "softbank",
        "samsung", "tsmc",
    ]),
    ("Big Political Events", [
        "election", "president", "presidential", "prime minister", "parliament", "congress",
        "senate", "senator", "white house", "kremlin", "sanction", "war", "ceasefire", "treaty",
        "summit", "nato", "european union", "united nations", "coup", "protest", "vote", "voter",
        "referendum", "minister", "ministry", "government", "trump", "putin", "zelensky",
        "xi jinping", "israel", "gaza", "palestine", "ukraine", "russia", "china", "iran",
        "tariff", "politic", "politics", "political", "politician", "democrat", "democracy",
        "republican", "labour", "tory", "tories", "conservative", "far right", "far left",
        "populist", "starmer", "macron", "merz", "modi", "erdogan", "netanyahu", "hamas",
        "houthi", "hezbollah", "taliban", "islamic state", "jihadist", "terror", "terrorism",
        "terrorist", "military", "army", "troops", "missile", "weapon", "nuclear", "immigration",
        "migrant", "asylum", "border", "deport", "supreme court", "court", "judge", "legislation",
        "lawmaker", "policy", "diplomat", "diplomatic", "diplomacy", "embassy", "regime",
        "dictator", "authoritarian", "opposition", "campaign", "poll", "mayor", "governor",
        "downing street", "brussels", "beijing", "moscow", "washington", "candidate", "party",
        "reform uk", "maga", "north korea", "taiwan", "pentagon", "cia", "fbi", "police", "crime",
        "gang", "mafia", "cartel", "corruption", "scandal", "impeach", "constitution", "invasion",
        "attack", "strike", "bomb", "conflict", "hostage", "refugee", "humanitarian", "brexit",
        "westminster", "capitol", "mp", "mps", "chancellor", "cabinet", "coalition",
        # "palestine" does not match "Palestinian" (the suffix rule only adds "ian" to the full
        # word), and "un" is needed for "UN meeting" / "the UN" (Le Monde visa story, Economist
        # secretary-general story)
        "palestinian", "un", "secretary general", "state department", "visa", "von der leyen",
        # crime / policing: Gladwell gun-violence piece, South African "maverick cop", Letby
        # inquiry. "police" does not match "policing".
        "gun", "shooting", "murder", "violence", "policing", "cop", "victim", "inquiry",
        "prison", "jail", "inmate", "citizen", "propaganda", "islamist", "islamism",
        "shadow chancellor", "shadow cabinet", "activist",
    ]),
    ("Companies & Finance", [
        "share", "shares", "stock", "stocks", "market", "markets", "earnings", "profit", "revenue",
        "ipo", "merger", "acquisition", "takeover", "deal", "investor", "invest", "investment",
        "bank", "banking", "fed", "federal reserve", "ecb", "bank of england", "boe", "boj",
        "bank of japan", "interest rate", "rate cut", "rate hike", "rate increase", "rate rise",
        "inflation", "gdp", "bond", "gilt", "dollar", "euro", "pound", "yen", "yuan", "oil", "ceo",
        "chief executive", "chairman", "board", "quarterly", "hedge fund", "private equity",
        "bankrupt", "bankruptcy", "layoff", "job cuts", "company", "companies", "firm", "business",
        "economy", "economic", "economics", "recession", "growth", "trade", "export", "import",
        "price", "cost", "wall street", "nasdaq", "s&p", "ftse", "fund", "asset", "debt", "loan",
        "mortgage", "pension", "tax", "budget", "wealth", "billionaire", "retail", "retailer",
        "consumer", "housing", "real estate", "property", "corporate", "industry", "manufacturing",
        "factory", "supply chain", "shipping", "airline", "automaker", "carmaker", "pharma",
        "drugmaker", "energy", "gas", "lng", "commodity", "commodities", "gold", "copper",
        "lithium", "mining", "monetary", "fiscal", "central bank", "treasury", "currency",
        "finance", "financial", "money", "wage", "salary", "pay", "jobs", "employment",
        "unemployment", "spending", "savings", "retirement", "brand", "sales", "customer",
        "shareholder", "buyout", "valuation", "tycoon", "luxury", "insurer", "insurance",
        "berkshire", "buffett", "glencore", "goldman", "jpmorgan", "blackrock", "boeing", "airbus",
        "toyota", "volkswagen", "shell", "exxon", "walmart", "hsbc", "barclays", "ubs", "nestle",
        "nestlé",
        # Economist "girl dad" managers piece and Ecuador prawn farmers piece; regulators probing
        # companies (Bloomberg travel-platform probe was landing in Culture via "travel"/"hotel")
        "manager", "hire", "hiring", "farm", "farmer", "agriculture", "probe", "antitrust",
        "regulator", "regulation",
    ]),
    ("Science & Health", [
        "scientist", "science", "research", "researcher", "study", "health", "doctor", "hospital",
        "disease", "cancer", "virus", "vaccine", "outbreak", "pandemic", "covid", "medicine",
        "medical", "drug", "brain", "gene", "genetic", "dna", "climate", "warming", "carbon",
        "emission", "weather", "flood", "drought", "wildfire", "hurricane", "earthquake",
        "glacier", "species", "wildlife", "animal", "bird", "ocean", "coral", "planet", "physics",
        "chemistry", "biology", "astronomy", "nasa", "mental health", "diet", "obesity", "sleep",
        "psychology", "measles", "fertility", "transplant", "surgery", "habitat", "soil",
        "environment", "environmental", "pollution", "nature", "forest", "river",
        "vitamin", "nurse", "nutrition",
    ]),
    ("Culture & Lifestyle", [
        "review", "film", "movie", "cinema", "director", "actor", "novel", "novelist", "book",
        "books", "author", "writer", "poet", "poetry", "fiction", "memoir", "art", "artist",
        "painting", "museum", "gallery", "exhibition", "music", "musician", "album", "song",
        "singer", "band", "concert", "orchestra", "opera", "ballet", "dance", "theatre", "theater",
        "broadway", "festival", "tv", "television", "series", "show", "documentary", "comedy",
        "comedian", "cartoon", "fashion", "style", "designer", "food", "restaurant", "chef",
        "recipe", "wine", "travel", "tourism", "tourist", "hotel", "holiday", "garden", "gardening",
        "wedding", "family", "parenting", "kids", "children", "teen", "teenager", "school",
        "college", "university", "student", "education", "teacher", "sport", "sports", "football",
        "soccer", "cricket", "tennis", "golf", "olympic", "nfl", "nba", "athlete", "celebrity",
        "royal", "monarchy", "prince", "king", "queen", "religion", "church", "history",
        "historian", "culture", "cultural", "lifestyle", "dating", "relationship", "podcast",
        "newsletter", "photograph", "photography", "architecture", "design", "literary", "prize",
        "award", "longlist", "obituary",
        # Obituaries: "dies aged 84", "has died". Bare "died"/"death" were here before but pulled
        # a Nigerian prison-deaths story into Culture, so only the obituary phrasings remain.
        "dies aged", "dies at", "has died",
        # Sport (Zidane squad story), music (A$AP Rocky), classroom/academic (New Yorker gym piece)
        "coach", "squad", "uefa", "world cup", "champions league", "premier league", "league",
        "rapper", "hip hop", "classroom", "academic",
    ]),
]

# Bonus points for the feed section an article came from: (section name -> {topic: points}).
# Section names are the keys used in news_rss.py FEEDS.
SECTION_BONUS = {
    "Technology": {"Technology & AI": 2},
    "Science & Technology": {"Science & Health": 2, "Technology & AI": 1},
    "Markets": {"Companies & Finance": 2},
    "Business": {"Companies & Finance": 2},
    "US Business": {"Companies & Finance": 2},
    "Companies": {"Companies & Finance": 2},
    "Economy": {"Companies & Finance": 2},
    "Economics": {"Companies & Finance": 2},
    "Finance & Economics": {"Companies & Finance": 2},
    "Personal Finance": {"Companies & Finance": 2},
    "Politics": {"Big Political Events": 2},
    "World": {"Big Political Events": 1},
    "World News": {"Big Political Events": 1},
    "International": {"Big Political Events": 1},
    "United States": {"Big Political Events": 1},
    "Europe": {"Big Political Events": 1},
    "Britain": {"Big Political Events": 1},
    "The Americas": {"Big Political Events": 1},
    "Middle East & Africa": {"Big Political Events": 1},
    "Asia": {"Big Political Events": 1},
    "China": {"Big Political Events": 1},
    "Leaders": {"Big Political Events": 1},
    # Semafor's regional / thematic categories (its single feed is split by <category>)
    "Gulf": {"Big Political Events": 1},
    "Africa": {"Big Political Events": 1},
    "Security": {"Big Political Events": 1},
    "Energy": {"Companies & Finance": 1},
    "Culture": {"Culture & Lifestyle": 2},
    "Arts & Culture": {"Culture & Lifestyle": 2},
    "Lifestyle": {"Culture & Lifestyle": 2},
    "Magazine": {"Culture & Lifestyle": 2},
}

# An article needs at least this score to get a topic, otherwise it is "Other".
MIN_TOPIC_SCORE = 2
OTHER_TOPIC = "Other"

# Educational / explainer detection works on the TITLE only.
EDUCATIONAL_PATTERNS = [
    # A question title: "Why Is It So Difficult to Regulate A.I.?"
    r"^\W*(how|why|what|when|where|who|which|is|are|should|could|can|does|do|will|would)\b.*\?\s*$",
    # Common explainer openings
    r"^\W*(how to|how the|how a|how an|why the|why a|what is|what are|what to know|what we know|what happens|what it means|what it takes|what you need)\b",
    # Explainer phrases anywhere in the title
    r"\b(explained|explainer|explains|a guide to|guide to|beginner'?s guide|what to know|what we know|need to know|the science of|the economics of|the history of|the case for|the case against|everything you need|q&a|in charts|in numbers|in data|the basics|lessons? from|understanding)\b",
    r"\bwhat (is|are|was|were|does|do|did|will|would|happens|happened)\b",
    r"\bwhy (is|are|was|were|does|do|did|the|it|we|you)\b",
    r"\bhow (to|does|do|did|can|could|will|would|the|a|an|much|many|far|long|it|we|you|i)\b",
]
# Titles with these words are never "educational" (reviews, podcasts, live blogs...)
NOT_EDUCATIONAL = r"\b(review|podcast|cartoon|newsletter|live|obituary|recipe)\b"

# Region tags. Add another region with one more entry.
REGIONS = {
    "Asia": {
        "keywords": [
            "asia", "asian", "china", "chinese", "beijing", "shanghai", "hong kong", "taiwan",
            "taipei", "japan", "japanese", "tokyo", "boj", "bank of japan", "korea", "korean",
            "seoul", "pyongyang", "kim jong", "india", "indian", "delhi", "mumbai", "modi",
            "pakistan", "bangladesh", "sri lanka", "nepal", "bhutan", "indonesia", "jakarta",
            "malaysia", "kuala lumpur", "singapore", "thailand", "bangkok", "vietnam", "hanoi",
            "philippines", "manila", "cambodia", "laos", "myanmar", "mongolia", "central asia",
            "southeast asia", "south asia", "east asia", "asia pacific", "asean", "xi jinping",
            "huawei", "tencent", "alibaba", "bytedance", "samsung", "toyota", "softbank", "tsmc",
            "nikkei", "yen", "yuan", "renminbi", "rupee", "himalaya",
        ],
        "sections": ["Asia", "China"],
    },
}

# --- Same-story grouping -----------------------------------------------------

# Words ignored when comparing titles (common words and news filler).
STOP_WORDS = set("""
the a an and or of to in on at for with from by as is are was were be been has have had this that
these those it its into over under after before about against amid among between during than then
their there they them his her he she we our you your who what when where why how not no new says
said say will would could should can may might more most much many some such also just only still
yet now up down out off all any one two three first last next big small long short high low old
young best worst good bad great little less least very really like get gets got make made take
took come came go went see saw know knew think thought want wants back way well even here time
year years day days week weeks month months today world people man men woman women life work home
city country state house us uk america american british french europe european review podcast
newsletter cartoon live latest another other becoming become plan plans right left ahead again
across along around behind beyond despite through toward without within enough every need needs
help helps keep keeps look looks give gives find finds turn turns call calls own same real true
never ever always often sometimes something anything nothing everything global
announce announces announced announcing administration market markets stock stocks since
highest lowest record records threat threats threatens threatened public service services
issue issues order orders ordered sign signs signed
trump trumps
""".split())
# The stop-word check runs on the raw word, before stemming, so each form is listed separately.
# "trump" is a stop word because it is in so many titles that it chained unrelated stories
# ("Oil Prices ... Trump's Iran War" + "Trump ... Medicaid Drug Prices"). Real Trump stories
# still group on their other words (greenland/security/deal, politico/barred/white).

# A shared title word only "counts" as strong evidence if it appears in at most this many titles.
RARE_MAX = 6
# Minimum shared title words for two articles to be the same story.
MIN_SHARED = 2
# ...unless at least one of them is rare, we need this many.
MIN_SHARED_WITHOUT_RARE = 3

# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------

_SUFFIX = r"(?:'s|s|es|ed|ing|d|er|ers|ist|ists|al|ic|ian|ians)?"


def normalize(text):
    """Lowercase and smooth out punctuation so keyword matching is predictable."""
    if not text:
        return ""
    text = text.lower()
    text = text.replace("’", "'").replace("‘", "'").replace("—", " ").replace("–", " ")
    text = re.sub(r"\b([a-z])\.([a-z])\.", r"\1\2", text)  # a.i. -> ai, u.s. -> us
    text = text.replace("-", " ")                            # data-centre -> data centre
    return text


def keyword_regex(keyword):
    """Regex for one keyword: whole words only, allowing common endings."""
    parts = [re.escape(p) for p in normalize(keyword).split()]
    body = r"\s+".join(parts)
    return re.compile(r"(?<![a-z])" + body + _SUFFIX + r"(?![a-z])")


def compile_keywords(keywords):
    return [(kw, keyword_regex(kw)) for kw in keywords]


def count_hits(compiled, text):
    """Return the keywords from `compiled` that appear in `text`."""
    return [kw for kw, rx in compiled if rx.search(text)]


_COMPILED_TOPICS = [(name, compile_keywords(kws)) for name, kws in TOPICS]
_COMPILED_REGIONS = {name: compile_keywords(cfg["keywords"]) for name, cfg in REGIONS.items()}
_EDU_PATTERNS = [re.compile(p) for p in EDUCATIONAL_PATTERNS]
_NOT_EDU = re.compile(NOT_EDUCATIONAL)

# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------


def load_articles(xlsx_path):
    """Read every sheet of the Excel file into a list of article dicts, without duplicates."""
    workbook = load_workbook(xlsx_path, read_only=True)
    articles = []
    seen_links = set()
    seen_titles = set()
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        next(rows, None)  # header row
        for row in rows:
            # Author is a 6th column added later; older spreadsheets simply lack it
            section, title, description, link, published, author = (list(row) + [None] * 6)[:6]
            title = (title or "").strip()
            link = (link or "").strip()
            if not title:
                continue
            title_key = (sheet.title, re.sub(r"\W+", " ", title.lower()).strip())
            if link in seen_links or title_key in seen_titles:
                continue
            seen_links.add(link)
            seen_titles.add(title_key)
            articles.append({
                "id": len(articles),
                "publication": sheet.title,
                "section": section or "",
                "title": title,
                "description": (description or "").strip(),
                "link": link,
                "published": published or "",
                "author": (author or "").strip(),
            })
    workbook.close()
    return articles

# ---------------------------------------------------------------------------
# Classification
# ---------------------------------------------------------------------------


def score_topics(article):
    """Score every topic: title hits count double, description hits once, plus section bonus."""
    title = normalize(article["title"])
    description = normalize(article["description"])
    scores = {}
    for name, compiled in _COMPILED_TOPICS:
        score = 2 * len(count_hits(compiled, title)) + len(count_hits(compiled, description))
        if score:
            scores[name] = score
    for name, bonus in SECTION_BONUS.get(article["section"], {}).items():
        scores[name] = scores.get(name, 0) + bonus
    return scores


def pick_topic(scores):
    if not scores:
        return OTHER_TOPIC
    best = max(scores, key=scores.get)
    if scores[best] < MIN_TOPIC_SCORE:
        return OTHER_TOPIC
    return best


def is_educational(title):
    title = normalize(title)
    if _NOT_EDU.search(title):
        return False
    for pattern in _EDU_PATTERNS:
        if pattern.search(title):
            return True
    return False


def matches_region(article, region_name):
    cfg = REGIONS[region_name]
    if article["section"] in cfg["sections"]:
        return True
    compiled = _COMPILED_REGIONS[region_name]
    if count_hits(compiled, normalize(article["title"])):
        return True
    return len(count_hits(compiled, normalize(article["description"]))) >= 2

# ---------------------------------------------------------------------------
# Same-story grouping
# ---------------------------------------------------------------------------


def stem(word):
    """Very small stemmer: rates -> rate, raises -> rais, elections -> election."""
    for suffix in ("ies", "ing", "ed", "es", "s"):
        if word.endswith(suffix) and len(word) - len(suffix) >= 4:
            if suffix == "ies":
                return word[:-3] + "y"
            return word[:-len(suffix)]
    return word


def title_tokens(title):
    """Return {stem: original word} for the meaningful words of a title."""
    text = title.lower().replace("’", "'")
    tokens = {}
    for word in re.findall(r"[^\W\d_][^\W\d_'\-]+", text):
        word = word.replace("'s", "").strip("'-")
        if len(word) >= 4 and word not in STOP_WORDS:
            tokens[stem(word)] = word
    return tokens


def group_same_story(articles):
    """Group articles from different publications whose titles share rare words.

    Uses "star" clustering: each group is one seed article plus everything related to
    that seed. This avoids chaining unrelated stories together (A~B, B~C but A!~C).
    """
    tokens = [set(title_tokens(a["title"])) for a in articles]
    doc_freq = collections.Counter(t for toks in tokens for t in toks)

    def related(i, j):
        if articles[i]["publication"] == articles[j]["publication"]:
            return False
        shared = tokens[i] & tokens[j]
        rare = [t for t in shared if doc_freq[t] <= RARE_MAX]
        if len(shared) >= MIN_SHARED and rare:
            return True
        return len(shared) >= MIN_SHARED_WITHOUT_RARE

    candidates = []
    for i in range(len(articles)):
        members = [i] + [j for j in range(len(articles)) if j != i and related(i, j)]
        publications = set(articles[k]["publication"] for k in members)
        if len(publications) >= 2:
            candidates.append(members)
    candidates.sort(key=len, reverse=True)

    used = set()
    groups = []
    for members in candidates:
        members = [k for k in members if k not in used]
        publications = set(articles[k]["publication"] for k in members)
        if len(publications) < 2:
            continue
        used.update(members)
        # Label: the words shared by at least two members, most common first
        shared_counter = collections.Counter(t for k in members for t in tokens[k])
        word_forms = {}
        for k in members:
            word_forms.update(title_tokens(articles[k]["title"]))
        label_words = [word_forms.get(t, t) for t, c in shared_counter.most_common(6) if c >= 2][:4]
        groups.append({
            "label": " · ".join(label_words),
            "article_ids": [articles[k]["id"] for k in members],
        })
    return groups

# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------


def build_report(xlsx_path):
    """Everything the web page needs, as one plain dict."""
    articles = load_articles(xlsx_path)
    counts = collections.Counter()
    for article in articles:
        scores = score_topics(article)
        article["topic"] = pick_topic(scores)
        article["educational"] = is_educational(article["title"])
        article["regions"] = [name for name in REGIONS if matches_region(article, name)]
        counts[article["topic"]] += 1
        if article["educational"]:
            counts["Educational"] += 1
        for region in article["regions"]:
            counts[region] += 1

    stories = group_same_story(articles)
    counts["Same story"] = len(stories)

    fetched_at = datetime.fromtimestamp(os.path.getmtime(xlsx_path)).strftime("%Y-%m-%d %H:%M")
    return {
        "ok": True,
        "fetched_at": fetched_at,
        "total": len(articles),
        "topics": [name for name, _ in TOPICS] + [OTHER_TOPIC],
        "regions": list(REGIONS),
        "publications": sorted(set(a["publication"] for a in articles)),
        "counts": dict(counts),
        "articles": articles,
        "stories": stories,
    }


if __name__ == "__main__":
    import json
    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "news_articles.xlsx")
    report = build_report(path)
    print(json.dumps(report["counts"], indent=2))
    print(f"{report['total']} articles, {len(report['stories'])} same-story groups")
